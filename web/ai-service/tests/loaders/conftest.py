"""Programmatic sample-file fixtures for loader tests.

Files are created fresh per test function so tests are completely isolated.
All fixture factories use the libraries already available as runtime deps
(pypdf, python-docx, openpyxl) plus fpdf2 (dev dep) for PDFs with text.
"""
import io

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF
from openpyxl import Workbook
from pypdf import PdfWriter


# ── PDF helpers ───────────────────────────────────────────────────────────────

def make_blank_pdf(pages: int = 1) -> bytes:
    """Valid PDF with `pages` blank pages (no extractable text)."""
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_text_pdf(text: str, pages: int = 1) -> bytes:
    """Valid PDF containing `text` on each page using fpdf2."""
    pdf = FPDF()
    for _ in range(pages):
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 10, text)
    return bytes(pdf.output())


@pytest.fixture
def valid_pdf_bytes() -> bytes:
    return make_text_pdf("Employee Handbook\nWelcome to Deloitte.", pages=2)


@pytest.fixture
def blank_pdf_bytes() -> bytes:
    return make_blank_pdf(pages=1)


@pytest.fixture
def corrupted_pdf_bytes() -> bytes:
    return b"%PDF-1.4 this is not a valid pdf structure"


@pytest.fixture
def empty_bytes() -> bytes:
    return b""


# ── DOCX helpers ──────────────────────────────────────────────────────────────

def make_docx(paragraphs: list[str], title: str = "", author: str = "") -> bytes:
    doc = DocxDocument()
    if title:
        doc.core_properties.title = title
    if author:
        doc.core_properties.author = author
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def valid_docx_bytes() -> bytes:
    return make_docx(
        ["Company Policy", "All employees must follow the code of conduct."],
        title="Company Policy",
        author="HR Department",
    )


@pytest.fixture
def corrupted_docx_bytes() -> bytes:
    return b"PK\x03\x04 not a real zip/docx"


# ── TXT helpers ───────────────────────────────────────────────────────────────

@pytest.fixture
def valid_txt_bytes() -> bytes:
    return (
        b"Onboarding Process\n"
        b"==================\n"
        b"Day 1: Meet your manager.\n"
        b"Day 2: Complete compliance training.\n"
        b"Day 3: Set up your workstation.\n"
    )


@pytest.fixture
def latin1_txt_bytes() -> bytes:
    return "Caf\xe9 au lait costs \xa35.00".encode("latin-1")


# ── CSV helpers ───────────────────────────────────────────────────────────────

@pytest.fixture
def valid_csv_bytes() -> bytes:
    return (
        b"name,department,leave_days\n"
        b"Alice,HR,20\n"
        b"Bob,IT,18\n"
        b"Carol,Finance,22\n"
    )


@pytest.fixture
def csv_with_bom_bytes() -> bytes:
    return "﻿name,department\nAlice,HR\n".encode("utf-8")


# ── XLSX helpers ─────────────────────────────────────────────────────────────

def make_xlsx(
    sheets: dict[str, list[list[str]]],
) -> bytes:
    """Create an XLSX with the given sheet data."""
    wb = Workbook()
    first = True
    for sheet_name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = sheet_name  # type: ignore[union-attr]
            first = False
        else:
            ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def valid_xlsx_bytes() -> bytes:
    return make_xlsx(
        {
            "Salary Grades": [
                ["Grade", "Min Salary", "Max Salary"],
                ["A", "30000", "50000"],
                ["B", "50000", "75000"],
                ["C", "75000", "110000"],
            ],
            "Holiday Calendar": [
                ["Date", "Holiday"],
                ["2026-01-01", "New Year's Day"],
                ["2026-12-25", "Christmas Day"],
            ],
        }
    )


@pytest.fixture
def corrupted_xlsx_bytes() -> bytes:
    return b"PK\x03\x04 not a real xlsx"
